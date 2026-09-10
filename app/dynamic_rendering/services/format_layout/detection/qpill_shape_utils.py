"""Locate question-pill slide shapes on input decks."""

from __future__ import annotations

from lxml import etree
from pptx.slide import Slide

from app.dynamic_rendering.utils.xml.helpers import local_name, prst_geom

INCH_EMU = 914_400
QUESTION_LABEL_TEXT = "Question"
OPTION_LABEL_TEXTS = frozenset({"A", "B", "C", "D"})
PILL_TO_TEXT_GAP_EMU = int(0.1 * INCH_EMU)
MCQ_ANSWER_TOP_TOLERANCE_EMU = int(1.5 * INCH_EMU)


def is_group(el: etree._Element) -> bool:
    return local_name(el) == "grpSp"


def _is_excluded_label(text: str) -> bool:
    return text == QUESTION_LABEL_TEXT or text in OPTION_LABEL_TEXTS


def question_pill_bottom_emu(slide: Slide) -> int | None:
    for sp in slide.shapes:
        if prst_geom(sp._element) == "roundRect" and sp.top is not None and sp.height is not None:
            return sp.top + sp.height
    return None


def first_mcq_option_top_emu(slide: Slide) -> int | None:
    tops: list[int] = []
    for sp in slide.shapes:
        elem = sp._element
        if is_group(elem) or prst_geom(elem) == "ellipse":
            if sp.top is not None:
                tops.append(sp.top)
    return min(tops) if tops else None


def find_question_text_element(
    slide: Slide,
    *,
    use_mcq_ceiling: bool = False,
) -> etree._Element | None:
    pill_bottom = question_pill_bottom_emu(slide)
    if pill_bottom is None:
        return None

    min_top = pill_bottom + PILL_TO_TEXT_GAP_EMU
    mcq_ceiling = first_mcq_option_top_emu(slide) if use_mcq_ceiling else None

    candidates: list[tuple[int, etree._Element]] = []
    for sp in slide.shapes:
        if not sp.has_text_frame or sp.top is None:
            continue
        text = sp.text_frame.text.strip()
        if len(text) < 1 or _is_excluded_label(text):
            continue
        if sp.top < min_top:
            continue
        if mcq_ceiling is not None and sp.top >= mcq_ceiling:
            continue
        candidates.append((sp.top, sp._element))

    if not candidates:
        return None
    candidates.sort(key=lambda item: item[0])
    return candidates[0][1]


def find_mcq_answer_text_elements(
    slide: Slide,
    question_text_el: etree._Element,
) -> list[etree._Element]:
    mcq_top = first_mcq_option_top_emu(slide)
    min_top = (mcq_top - MCQ_ANSWER_TOP_TOLERANCE_EMU) if mcq_top is not None else 0

    candidates: list[tuple[int, etree._Element]] = []
    for sp in slide.shapes:
        if not sp.has_text_frame or sp.top is None:
            continue
        elem = sp._element
        if elem is question_text_el:
            continue
        text = sp.text_frame.text.strip()
        if len(text) < 1 or _is_excluded_label(text):
            continue
        if sp.top < min_top:
            continue
        candidates.append((sp.top, elem))

    candidates.sort(key=lambda item: item[0])
    return [el for _, el in candidates[:4]]


def has_question_text_below_pill(slide: Slide) -> bool:
    return find_question_text_element(slide, use_mcq_ceiling=False) is not None

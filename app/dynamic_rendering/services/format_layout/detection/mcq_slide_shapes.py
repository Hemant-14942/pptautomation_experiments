"""Shared shape discovery for question-pill + MCQ input slides."""

from __future__ import annotations

from pptx.slide import Slide

from app.dynamic_rendering.services.format_layout.detection.input_mcq_options import (
    collect_input_mcq_options,
)
from app.dynamic_rendering.services.format_layout.detection.qpill_shape_utils import (
    find_mcq_answer_text_elements,
    find_question_text_element,
)
from app.dynamic_rendering.utils.xml.helpers import prst_geom


def find_mcq_slide_shapes(slide: Slide) -> dict | None:
    question_pill_el = None
    question_label_el = None

    for sp in slide.shapes:
        elem = sp._element
        geom = prst_geom(elem)
        if geom == "roundRect" and question_pill_el is None:
            question_pill_el = elem
            continue
        if not sp.has_text_frame:
            continue
        text = sp.text_frame.text.strip()
        if text == "Question" and question_label_el is None:
            question_label_el = elem

    options = collect_input_mcq_options(slide)
    question_text_el = find_question_text_element(slide, use_mcq_ceiling=True)
    mcq_text_els = (
        find_mcq_answer_text_elements(slide, question_text_el, limit=len(options))
        if question_text_el is not None
        else []
    )

    if not (
        question_pill_el is not None
        and question_label_el is not None
        and question_text_el is not None
        and len(options) >= 1
    ):
        return None

    # Pair pills top-to-bottom with answer texts in the same order.
    # Missing answer text for a pill is fine — still format the pill layout slot.
    mcq_options = []
    for option_idx, opt in enumerate(options):
        text_el = mcq_text_els[option_idx] if option_idx < len(mcq_text_els) else None
        mcq_options.append(
            {
                "pill_el": opt["pill_el"],
                "label_el": opt["label_el"],
                "letter": opt["letter"],
                "option_idx": option_idx,
                "text_el": text_el,
                "grouped": opt["grouped"],
            }
        )

    return {
        "question_pill_el": question_pill_el,
        "question_label_el": question_label_el,
        "question_text_el": question_text_el,
        "mcq_options": mcq_options,
    }

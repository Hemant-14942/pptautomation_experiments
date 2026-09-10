from typing import Any

from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide

from app.dynamic_rendering.constants.xml_namespaces import NS
from app.dynamic_rendering.services.format_layout.detection.qpill_shape_utils import (
    has_question_text_below_pill,
)
from app.dynamic_rendering.services.text.heading_detector import extract_text_from_slide

SLIDE_TYPE_TITLE_TABLE_ONLY = "title_table_only"
SLIDE_TYPE_TITLE_BODY_ONLY = "title_body_only"
SLIDE_TYPE_TITLE_BODY_SINGLE_IMAGE = "title_body_single_image"
SLIDE_TYPE_TITLE_BODY_MULTIPLE_IMAGES = "title_body_multiple_images"
SLIDE_TYPE_TITLE_IMAGE_ONLY = "title_image_only"

SLIDE_TYPE_BODY_ONLY = "body_only"
SLIDE_TYPE_IMAGE_ONLY = "image_only"
SLIDE_TYPE_TABLE_ONLY = "table_only"

SLIDE_TYPE_QUESTION_QTEXT_MCQ = "question_qtext_mcq"
SLIDE_TYPE_QPILL_QTEXT_TABLE_MCQ = "qpill_qtext_table_mcq"
SLIDE_TYPE_QPILL_QTEXT_ONLY = "qpill_qtext_only"
SLIDE_TYPE_QPILL_QTEXT_IMAGE_MCQ = "qpill_qtext_image_mcq"

IMAGE_Y_THRESHOLD_EMU = 1_828_800  # 2.0 inches


def has_real_body_text(slide: Slide, heading_pos: tuple[int, int] | None) -> bool:
    for sp in slide.shapes:
        if not sp.has_text_frame:
            continue
        if heading_pos is not None and (sp.left, sp.top) == heading_pos:
            continue
        if len(sp.text_frame.text.strip()) > 25:
            return True
    return False


def count_real_images(slide: Slide) -> int:
    count = 0
    for sp in slide.shapes:
        shape_type = sp.shape_type
        if shape_type is not None and getattr(shape_type, "name", "") == "PICTURE":
            if sp.top is not None and sp.top > IMAGE_Y_THRESHOLD_EMU:
                count += 1
    return count


def count_actual_tables(slide: Slide) -> int:
    table_count = 0
    for sp in slide.shapes:
        shape_type = sp.shape_type
        if shape_type is not None and getattr(shape_type, "name", "") == "TABLE":
            table_count += 1
    return table_count


def has_question_pill(slide: Slide) -> bool:
    for sp in slide.shapes:
        if not hasattr(sp, "_element"):
            continue
        geom = sp._element.find(".//{%s}prstGeom" % NS["a"])
        if geom is not None and geom.get("prst") == "roundRect":
            return True
    return False


def has_option_pill_shapes(slide: Slide) -> bool:
    for sp in slide.shapes:
        if str(sp.shape_type) == "<ShapeType.GROUP: 6>":
            return True
        if not hasattr(sp, "_element"):
            continue
        geom = sp._element.find(".//{%s}prstGeom" % NS["a"])
        if geom is not None and geom.get("prst") == "ellipse":
            return True
    return False


def has_any_text_content(slide: Slide) -> bool:
    for sp in slide.shapes:
        if sp.has_text_frame and sp.text_frame.text.strip():
            return True
    return False


def _no_title_or_question_layout(
    slide: Slide, prs: PresentationType, signature: dict[str, Any]
) -> bool:
    if extract_text_from_slide(slide, 0, prs).get("heading") is not None:
        return False
    if signature.get("option_pills", 0) > 0:
        return False
    if has_question_pill(slide):
        return False
    if has_option_pill_shapes(slide):
        return False
    if has_mcq_structure(slide):
        return False
    return True


def qualifies_for_pure_slide(slide: Slide, prs: PresentationType, signature: dict[str, Any]) -> bool:
    if not _no_title_or_question_layout(slide, prs, signature):
        return False
    if count_actual_tables(slide) > 0:
        return False
    return True


def qualifies_for_table_only_slide(
    slide: Slide, prs: PresentationType, signature: dict[str, Any]
) -> bool:
    if not _no_title_or_question_layout(slide, prs, signature):
        return False
    if count_real_images(slide) > 0:
        return False
    if has_any_text_content(slide):
        return False
    return count_actual_tables(slide) == 1


def has_mcq_structure(slide: Slide) -> bool:
    roundrect_count = 0
    ellipse_count = 0

    for sp in slide.shapes:
        if hasattr(sp, "_element"):
            elem = sp._element
            geom = elem.find(".//{%s}prstGeom" % NS["a"])
            if geom is not None:
                prst = geom.get("prst")
                if prst == "roundRect":
                    roundrect_count += 1
                elif prst == "ellipse":
                    ellipse_count += 1

        if str(sp.shape_type) == "<ShapeType.GROUP: 6>":
            ellipse_count += 1

    return roundrect_count >= 1 and ellipse_count >= 4


def qualifies_for_qpill_qtext_only_slide(
    slide: Slide, prs: PresentationType, signature: dict[str, Any]
) -> bool:
    if extract_text_from_slide(slide, 0, prs).get("heading") is not None:
        return False
    if not has_question_pill(slide):
        return False
    if has_mcq_structure(slide) or has_option_pill_shapes(slide):
        return False
    if count_actual_tables(slide) > 0:
        return False
    if count_real_images(slide) > 0:
        return False
    return has_question_text_below_pill(slide)


def detect_slide_type(
    signature: dict[str, Any], slide: Slide, prs: PresentationType
) -> str | None:
    if has_mcq_structure(slide):
        actual_table_count = count_actual_tables(slide)
        real_image_count = count_real_images(slide)

        if actual_table_count == 0 and real_image_count >= 1:
            return SLIDE_TYPE_QPILL_QTEXT_IMAGE_MCQ
        if real_image_count == 0:
            if actual_table_count == 1:
                return SLIDE_TYPE_QPILL_QTEXT_TABLE_MCQ
            if actual_table_count == 0:
                return SLIDE_TYPE_QUESTION_QTEXT_MCQ

    if qualifies_for_qpill_qtext_only_slide(slide, prs, signature):
        return SLIDE_TYPE_QPILL_QTEXT_ONLY

    heading = extract_text_from_slide(slide, 0, prs).get("heading")
    has_heading = heading is not None
    no_options = signature.get("option_pills", 0) == 0
    table_count = signature.get("table_count", 0)
    real_image_count = count_real_images(slide)

    if qualifies_for_table_only_slide(slide, prs, signature):
        return SLIDE_TYPE_TABLE_ONLY

    if qualifies_for_pure_slide(slide, prs, signature):
        if real_image_count >= 1 and not has_any_text_content(slide):
            return SLIDE_TYPE_IMAGE_ONLY
        if real_image_count == 0 and has_real_body_text(slide, None):
            return SLIDE_TYPE_BODY_ONLY

    if not (has_heading and no_options):
        return None

    heading_pos = heading.get("position_top_left")
    has_body = has_real_body_text(slide, heading_pos)

    if table_count == 1 and not has_body and real_image_count == 0:
        return SLIDE_TYPE_TITLE_TABLE_ONLY

    if table_count == 0 and not has_body and real_image_count >= 1:
        return SLIDE_TYPE_TITLE_IMAGE_ONLY

    if table_count == 0 and has_body:
        if real_image_count == 0:
            return SLIDE_TYPE_TITLE_BODY_ONLY
        if real_image_count == 1:
            return SLIDE_TYPE_TITLE_BODY_SINGLE_IMAGE
        if real_image_count >= 2:
            return SLIDE_TYPE_TITLE_BODY_MULTIPLE_IMAGES

    return None

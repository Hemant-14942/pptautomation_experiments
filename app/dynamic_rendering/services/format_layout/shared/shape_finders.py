"""Shared shape-finding helpers for fit-layout formatters."""

from __future__ import annotations

from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide

from app.dynamic_rendering.services.format_layout.detection.slide_type_detector import IMAGE_Y_THRESHOLD_EMU
from app.dynamic_rendering.services.text.heading_detector import extract_text_from_slide
from app.dynamic_rendering.utils.xml.helpers import local_name


def find_body_el(slide: Slide, prs: PresentationType | None = None) -> object | None:
    """Body text shape (>25 chars), excluding the title heading shape."""
    heading_pos = None
    if prs is not None:
        heading = extract_text_from_slide(slide, 0, prs).get("heading")
        heading_pos = heading.get("position_top_left") if heading else None

    for sp in slide.shapes:
        if not sp.has_text_frame:
            continue
        if heading_pos is not None and (sp.left, sp.top) == heading_pos:
            continue
        if len(sp.text_frame.text.strip()) > 25:
            return sp._element
    return None


def find_body_only_el(slide: Slide) -> object | None:
    best_el = None
    best_len = 0
    for sp in slide.shapes:
        if not sp.has_text_frame:
            continue
        text_len = len(sp.text_frame.text.strip())
        if text_len > 25 and text_len > best_len:
            best_len = text_len
            best_el = sp._element
    return best_el


def find_table_el(slide: Slide) -> object | None:
    for sp in slide.shapes:
        if sp.has_table:
            return sp._element
    return None


def find_content_pictures(slide: Slide) -> list[tuple[object, tuple[int, int]]]:
    """Return [(picture_el, (width_px, height_px)), ...] below header area."""
    pics: list[tuple[object, tuple[int, int]]] = []
    for sp in slide.shapes:
        shape_type = sp.shape_type
        if shape_type is not None and getattr(shape_type, "name", "") == "PICTURE":
            if sp.top is not None and sp.top > IMAGE_Y_THRESHOLD_EMU:
                pics.append((sp._element, sp.image.size))
    return pics

"""Locate shapes on a qpill_qtext_only input slide."""

from __future__ import annotations

from pptx.slide import Slide

from app.dynamic_rendering.services.format_layout.detection.qpill_shape_utils import (
    find_question_text_element,
)
from app.dynamic_rendering.utils.xml.helpers import prst_geom


def find_shapes(slide: Slide) -> dict | None:
    question_pill_el = None
    question_label_el = None

    for sp in slide.shapes:
        elem = sp._element
        if prst_geom(elem) == "roundRect" and question_pill_el is None:
            question_pill_el = elem
            continue

        if not sp.has_text_frame:
            continue
        text = sp.text_frame.text.strip()
        if not text:
            continue
        if text == "Question" and question_label_el is None:
            question_label_el = elem

    question_text_el = find_question_text_element(slide, use_mcq_ceiling=False)

    if question_pill_el is None or question_label_el is None or question_text_el is None:
        return None

    return {
        "question_pill_el": question_pill_el,
        "question_label_el": question_label_el,
        "question_text_el": question_text_el,
    }

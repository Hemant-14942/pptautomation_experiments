"""Locate shapes on a qpill_qtext_multiple_image input slide."""

from __future__ import annotations

from pptx.slide import Slide

from app.dynamic_rendering.services.format_layout.detection.qpill_shape_utils import (
    find_question_text_element,
)
from app.dynamic_rendering.services.format_layout.detection.slide_type_detector import (
    IMAGE_Y_THRESHOLD_EMU,
)
from app.dynamic_rendering.utils.xml.helpers import prst_geom


def find_shapes(slide: Slide) -> dict | None:
    question_pill_el = None
    question_label_el = None

    for sp in slide.shapes:
        elem = sp._element
        if prst_geom(elem) == "roundRect" and question_pill_el is None:
            question_pill_el = elem
            continue

        if not sp.has_text_frame:
            continue
        text = sp.text_frame.text.strip()
        if not text:
            continue
        if text == "Question" and question_label_el is None:
            question_label_el = elem

    question_text_el = find_question_text_element(slide, use_mcq_ceiling=False)

    pictures: list[tuple[int, int, object, tuple[int, int]]] = []
    for sp in slide.shapes:
        shape_type = sp.shape_type
        if shape_type is not None and getattr(shape_type, "name", "") == "PICTURE":
            if sp.top is not None and sp.top > IMAGE_Y_THRESHOLD_EMU:
                pictures.append((sp.top, sp.left, sp._element, sp.image.size))

    if question_pill_el is None or question_label_el is None or question_text_el is None:
        return None
    if len(pictures) < 2:
        return None

    pictures.sort(key=lambda item: (item[0], item[1]))

    return {
        "question_pill_el": question_pill_el,
        "question_label_el": question_label_el,
        "question_text_el": question_text_el,
        "picture_els": [p[2] for p in pictures],
        "picture_sizes": [p[3] for p in pictures],
    }

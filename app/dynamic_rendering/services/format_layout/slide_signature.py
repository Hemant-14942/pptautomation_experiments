"""Lightweight per-slide structural signature for layout type detection."""

from __future__ import annotations

from typing import Any

from pptx import Presentation


def parse_input_slide_signature(input_path: str) -> list[dict[str, Any]]:
    """Count shapes per slide — used as hints by detect_slide_type."""
    prs = Presentation(input_path)
    sigs: list[dict[str, Any]] = []
    for idx, slide in enumerate(prs.slides):
        text_count = 0
        group_count = 0
        table_count = 0
        picture_count = 0
        option_pill_count = 0
        question_badge = False
        body_text_count = 0
        sample = ""
        for sp in slide.shapes:
            tag = sp.shape_type
            if tag is None:
                group_count += 1
                try:
                    cw = sp.width or 0
                    ch = sp.height or 0
                    if cw and ch and cw < 2000000 and ch < 2000000:
                        option_pill_count += 1
                except Exception:
                    pass
                continue
            if hasattr(tag, "name") and tag.name == "TABLE":
                table_count += 1
            elif hasattr(tag, "name") and tag.name == "PICTURE":
                picture_count += 1
            elif sp.has_text_frame:
                text_count += 1
                txt = sp.text_frame.text.strip().lower()
                if not sample:
                    sample = sp.text_frame.text.strip()[:120]
                if "question" in txt or "questions" in txt:
                    question_badge = True
                if len(txt) > 30:
                    body_text_count += 1
        sigs.append(
            {
                "index": idx,
                "shape_count": len(slide.shapes),
                "text_count": text_count,
                "group_count": group_count,
                "table_count": table_count,
                "picture_count": picture_count,
                "option_pills": option_pill_count,
                "has_question_badge": question_badge,
                "body_text_count": body_text_count,
                "sample": sample,
            }
        )
    return sigs

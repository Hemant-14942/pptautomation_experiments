"""Test the body_only formatter.

Run on one slide manually:
    python -m app.core.fit_layout.layout_formatter.body_only.tests.test_formatter <pptx_path> <slide_index>
"""

import os
import sys

from pptx import Presentation
from pptx.util import Emu

from app.core.slide_text_analyzer import extract_text_from_slide
from app.core.style_parser import parse_input_slide_signature
from app.core.fit_layout.layout_formatter.slide_type_detector import (
    SLIDE_TYPE_BODY_ONLY,
    detect_slide_type,
)
from app.core.fit_layout.layout_formatter.body_only.constants import BODY_TEXTBOX
from app.core.fit_layout.layout_formatter.body_only.formatter import format_body_text

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "output")

XFRM_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def _get_off_ext(el):
    """Read position (off) and size (ext) from the shape's XML."""
    xfrm = el.find(".//p:spPr/a:xfrm", XFRM_NS)
    off = xfrm.find("a:off", XFRM_NS)
    ext = xfrm.find("a:ext", XFRM_NS)
    return (int(off.get("x")), int(off.get("y"))), (int(ext.get("cx")), int(ext.get("cy")))


def _find_body_el(slide):
    """Find the main body text shape — the one with the most text content.

    On a body_only slide there is no heading, so we just pick the text
    shape that has the longest text (more than 30 characters).
    """
    best_el = None
    best_len = 0
    for sp in slide.shapes:
        if not sp.has_text_frame:
            continue
        text_len = len(sp.text_frame.text.strip())
        if text_len > 30 and text_len > best_len:
            best_len = text_len
            best_el = sp._element
    return best_el


def _get_anchor(el):
    """Read the vertical anchor value from bodyPr (should be 'ctr')."""
    txBody = el.find(".//p:txBody", XFRM_NS)
    if txBody is None:
        return None
    bodyPr = txBody.find("a:bodyPr", XFRM_NS)
    if bodyPr is None:
        return None
    return bodyPr.get("anchor")


def _has_norm_autofit(el):
    """Check that normAutofit tag exists (shrink text on overflow)."""
    txBody = el.find(".//p:txBody", XFRM_NS)
    if txBody is None:
        return False
    bodyPr = txBody.find("a:bodyPr", XFRM_NS)
    if bodyPr is None:
        return False
    autofit = bodyPr.find("a:normAutofit", XFRM_NS)
    return autofit is not None


def verify_slide(pptx_path: str, slide_index: int) -> dict:
    """Check one slide: does the detector say body_only, and does the
    formatter place the text box in the right spot with the right settings?"""
    prs = Presentation(pptx_path)
    slide = prs.slides[slide_index]
    sig = parse_input_slide_signature(pptx_path)[slide_index]

    detected_type = detect_slide_type(sig, slide, prs)
    if detected_type != SLIDE_TYPE_BODY_ONLY:
        return {
            "pass": False,
            "reason": f"slide type detected as {detected_type!r}, not body_only",
            "detected_type": detected_type,
        }

    body_el = _find_body_el(slide)
    if body_el is None:
        return {
            "pass": False,
            "reason": "detector matched but could not locate body text shape",
            "detected_type": detected_type,
        }

    body_clone = format_body_text(body_el)
    body_off, body_ext = _get_off_ext(body_clone)
    anchor = _get_anchor(body_clone)
    autofit_ok = _has_norm_autofit(body_clone)

    body_ok = (body_off, body_ext) == (
        (BODY_TEXTBOX["x"], BODY_TEXTBOX["y"]),
        (BODY_TEXTBOX["width"], BODY_TEXTBOX["height"]),
    )
    anchor_ok = anchor == "ctr"

    return {
        "pass": body_ok and anchor_ok and autofit_ok,
        "detected_type": detected_type,
        "body": {
            "off": body_off,
            "ext": body_ext,
            "matches_fixed_box": body_ok,
        },
        "anchor": {"value": anchor, "is_center": anchor_ok},
        "norm_autofit": autofit_ok,
    }


def save_formatted_slide(pptx_path: str, slide_index: int, output_dir: str = OUTPUT_DIR) -> str | None:
    """Apply the formatter to one slide and save a new pptx file."""
    prs = Presentation(pptx_path)
    slide = prs.slides[slide_index]
    sig = parse_input_slide_signature(pptx_path)[slide_index]

    if detect_slide_type(sig, slide, prs) != SLIDE_TYPE_BODY_ONLY:
        return None

    body_el = _find_body_el(slide)
    if body_el is None:
        return None

    body_clone = format_body_text(body_el)
    spTree = slide.shapes._spTree
    spTree.remove(body_el)
    spTree.append(body_clone)

    os.makedirs(output_dir, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(pptx_path))[0]
    out_path = os.path.join(output_dir, f"{base_name}_slide{slide_index + 1}_formatted.pptx")
    prs.save(out_path)
    return out_path


def _print_report(pptx_path: str, slide_index: int, report: dict) -> None:
    print(f"\n{pptx_path} | slide index {slide_index} (PPT slide {slide_index + 1})")
    print(f"  detected_type: {report.get('detected_type')}")
    if not report["pass"] and "reason" in report:
        print(f"  RESULT: FAIL — {report['reason']}")
        return

    b = report["body"]
    a = report["anchor"]
    b_off_in = (Emu(b["off"][0]).inches, Emu(b["off"][1]).inches)
    b_ext_in = (Emu(b["ext"][0]).inches, Emu(b["ext"][1]).inches)

    print(f"  body:   off={b_off_in[0]:.2f},{b_off_in[1]:.2f}in  ext={b_ext_in[0]:.2f}x{b_ext_in[1]:.2f}in  fixed_box_match={b['matches_fixed_box']}")
    print(f"  anchor: {a['value']!r}  center={a['is_center']}")
    print(f"  normAutofit: {report.get('norm_autofit')}")
    print(f"  RESULT: {'PASS' if report['pass'] else 'FAIL'}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python -m ...test_formatter <pptx_path> <slide_index>")
        sys.exit(1)
    path, idx = sys.argv[1], int(sys.argv[2])
    result = verify_slide(path, idx)
    _print_report(path, idx, result)
    if result["pass"]:
        saved_path = save_formatted_slide(path, idx)
        print(f"  saved formatted pptx: {saved_path}")

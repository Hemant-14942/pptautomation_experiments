"""Test title_image_only formatter.

Run one slide:
    python -m app.core.fit_layout.layout_formatter.title_image_only.tests.test_formatter app/data/input/ptest1.pptx 17

Run all matching slides:
    python -m app.core.fit_layout.layout_formatter.title_image_only.tests.test_formatter app/data/input/ptest1.pptx
"""

import os
import sys

from pptx import Presentation

from app.core.style_parser import parse_input_slide_signature
from app.core.fit_layout.layout_formatter.slide_type_detector import (
    IMAGE_Y_THRESHOLD_EMU,
    SLIDE_TYPE_TITLE_IMAGE_ONLY,
    detect_slide_type,
)
from app.core.fit_layout.layout_formatter.title_image_only.constants import get_title_image_boxes
from app.core.fit_layout.layout_formatter.title_image_only.formatter import format_title_image_only_slide

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "output")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def _get_off_ext(el):
    xfrm = el.find(".//p:spPr/a:xfrm", NS)
    off = xfrm.find("a:off", NS)
    ext = xfrm.find("a:ext", NS)
    return (
        (int(float(off.get("x"))), int(float(off.get("y")))),
        (int(float(ext.get("cx"))), int(float(ext.get("cy")))),
    )


def _find_content_pictures(slide):
    pictures = []
    for sp in slide.shapes:
        stype = sp.shape_type
        if stype is not None and getattr(stype, "name", "") == "PICTURE":
            if sp.top is not None and sp.top > IMAGE_Y_THRESHOLD_EMU:
                pictures.append((sp.top, sp.left, sp._element, sp.image.size))
    pictures.sort(key=lambda item: (item[0], item[1]))
    return [p[2] for p in pictures], [p[3] for p in pictures]


def save_formatted_slide(pptx_path: str, slide_index: int, output_dir: str = OUTPUT_DIR) -> str | None:
    prs = Presentation(pptx_path)
    slide = prs.slides[slide_index]
    sig = parse_input_slide_signature(pptx_path)[slide_index]

    if detect_slide_type(sig, slide, prs) != SLIDE_TYPE_TITLE_IMAGE_ONLY:
        return None

    pic_els, pic_sizes = _find_content_pictures(slide)
    if not pic_els:
        return None

    clones = format_title_image_only_slide(pic_els, pic_sizes)
    spTree = slide.shapes._spTree
    for orig_el, clone_el in zip(pic_els, clones):
        if orig_el.getparent() is not None:
            spTree.remove(orig_el)
        spTree.append(clone_el)

    os.makedirs(output_dir, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(pptx_path))[0]
    out_path = os.path.join(output_dir, f"{base_name}_slide{slide_index + 1}_formatted.pptx")
    prs.save(out_path)
    return out_path


def run_all(pptx_path: str) -> None:
    prs = Presentation(pptx_path)
    signatures = parse_input_slide_signature(pptx_path)

    matched = 0
    saved = 0
    for idx, slide in enumerate(prs.slides):
        if detect_slide_type(signatures[idx], slide, prs) != SLIDE_TYPE_TITLE_IMAGE_ONLY:
            continue
        matched += 1
        pic_els, _ = _find_content_pictures(slide)
        boxes = get_title_image_boxes(len(pic_els))
        out_path = save_formatted_slide(pptx_path, idx)
        status = out_path if out_path else "SKIPPED"
        print(f"  slide {idx + 1}: {len(pic_els)} images, layout={len(boxes)} slots -> {status}")
        if out_path:
            saved += 1

    print(f"\n{matched} slides detected as title_image_only, {saved} saved to {OUTPUT_DIR}")


if __name__ == "__main__":
    if len(sys.argv) == 3:
        path, idx = sys.argv[1], int(sys.argv[2])
        prs = Presentation(path)
        sig = parse_input_slide_signature(path)[idx]
        print(f"detected: {detect_slide_type(sig, prs.slides[idx], prs)}")
        out = save_formatted_slide(path, idx)
        print(f"saved: {out}" if out else "no match / no pictures")
    elif len(sys.argv) == 2:
        run_all(sys.argv[1])
    else:
        print("Usage: python -m ...test_formatter <pptx_path> [<slide_index>]")
        sys.exit(1)

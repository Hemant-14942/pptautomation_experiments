"""Test qpill_qtext_table_mcq formatter.

Run:
    python -m app.core.fit_layout.layout_formatter.qpill_qtext_table_mcq.tests.test_formatter app/data/input/p.pptx 0
"""

import os
import sys

from pptx import Presentation

from app.core.style_parser import parse_input_slide_signature
from app.core.fit_layout.layout_formatter.slide_type_detector import (
    SLIDE_TYPE_QPILL_QTEXT_TABLE_MCQ,
    detect_slide_type,
)
from app.core.fit_layout.layout_formatter.qpill_qtext_mcq.formatter import is_group
from app.core.fit_layout.layout_formatter.qpill_shape_utils import (
    find_mcq_answer_text_elements,
    find_question_text_element,
)
from app.core.fit_layout.layout_formatter.qpill_qtext_table_mcq.formatter import (
    format_mcq_answer_text,
    format_mcq_option_label,
    format_mcq_option_pill,
    format_question_label,
    format_question_pill,
    format_question_text,
    format_table,
)
from app.utils.xml_helpers import prst_geom

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "output")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def _has_norm_autofit(el) -> bool:
    txBody = el.find(".//p:txBody", NS)
    if txBody is None:
        return False
    bodyPr = txBody.find("a:bodyPr", NS)
    return bodyPr is not None and bodyPr.find("a:normAutofit", NS) is not None


def find_shapes(slide) -> dict | None:
    question_pill_el = None
    question_label_el = None
    question_text_el = None
    table_el = None
    mcq_pill_els = []
    mcq_label_els = {}
    mcq_text_els = []

    for sp in slide.shapes:
        elem = sp._element

        if sp.has_table:
            table_el = elem
            continue

        if is_group(elem):
            mcq_pill_els.append(elem)
            continue

        geom = prst_geom(elem)
        if geom == "roundRect" and question_pill_el is None:
            question_pill_el = elem
            continue
        if geom == "ellipse":
            mcq_pill_els.append(elem)
            continue

        if not sp.has_text_frame:
            continue
        text = sp.text_frame.text.strip()
        if not text:
            continue

        if text == "Question" and question_label_el is None:
            question_label_el = elem
        elif text in ("A", "B", "C", "D") and text not in mcq_label_els:
            mcq_label_els[text] = elem

    question_text_el = find_question_text_element(slide, use_mcq_ceiling=True)
    if question_text_el is not None:
        mcq_text_els = find_mcq_answer_text_elements(slide, question_text_el)

    if not (
        question_pill_el is not None
        and question_label_el is not None
        and question_text_el is not None
        and table_el is not None
        and len(mcq_pill_els) >= 4
        and len(mcq_text_els) >= 4
    ):
        return None

    return {
        "question_pill_el": question_pill_el,
        "question_label_el": question_label_el,
        "question_text_el": question_text_el,
        "table_el": table_el,
        "mcq_pill_els": mcq_pill_els[:4],
        "mcq_label_els": mcq_label_els,
        "mcq_text_els": mcq_text_els[:4],
    }


def save_formatted_slide(pptx_path: str, slide_index: int, output_dir: str = OUTPUT_DIR) -> str | None:
    prs = Presentation(pptx_path)
    slide = prs.slides[slide_index]
    sig = parse_input_slide_signature(pptx_path)[slide_index]

    if detect_slide_type(sig, slide, prs) != SLIDE_TYPE_QPILL_QTEXT_TABLE_MCQ:
        return None

    shapes = find_shapes(slide)
    if shapes is None:
        return None

    spTree = slide.shapes._spTree

    def reposition(orig_el, clone_el):
        if orig_el.getparent() is not None:
            spTree.remove(orig_el)
        spTree.append(clone_el)

    reposition(shapes["question_pill_el"], format_question_pill(shapes["question_pill_el"]))
    reposition(shapes["question_label_el"], format_question_label(shapes["question_label_el"]))
    reposition(shapes["question_text_el"], format_question_text(shapes["question_text_el"]))
    reposition(shapes["table_el"], format_table(shapes["table_el"]))

    letters = ["A", "B", "C", "D"]
    for opt_idx, pill_el in enumerate(shapes["mcq_pill_els"]):
        clone = format_mcq_option_pill(opt_idx, pill_el)
        if clone is not None:
            reposition(pill_el, clone)
            label_el = shapes["mcq_label_els"].get(letters[opt_idx])
            if label_el is not None:
                reposition(label_el, format_mcq_option_label(opt_idx, label_el))
        reposition(shapes["mcq_text_els"][opt_idx], format_mcq_answer_text(opt_idx, shapes["mcq_text_els"][opt_idx]))

    os.makedirs(output_dir, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(pptx_path))[0]
    out_path = os.path.join(output_dir, f"{base_name}_slide{slide_index + 1}_formatted.pptx")
    prs.save(out_path)
    return out_path


def run_all(pptx_path: str) -> None:
    prs = Presentation(pptx_path)
    signatures = parse_input_slide_signature(pptx_path)

    matched = 0
    saved = 0
    for idx, slide in enumerate(prs.slides):
        if detect_slide_type(signatures[idx], slide, prs) != SLIDE_TYPE_QPILL_QTEXT_TABLE_MCQ:
            continue
        matched += 1
        out_path = save_formatted_slide(pptx_path, idx)
        status = out_path if out_path else "SKIPPED (shapes not found)"
        print(f"  slide index {idx:2d} (PPT slide {idx + 1:2d}): {status}")
        if out_path:
            saved += 1

    print(f"\n{matched} slides detected as qpill_qtext_table_mcq, {saved} formatted and saved to {OUTPUT_DIR}")


if __name__ == "__main__":
    if len(sys.argv) == 3:
        path, idx = sys.argv[1], int(sys.argv[2])
        prs = Presentation(path)
        sig = parse_input_slide_signature(path)[idx]
        detected = detect_slide_type(sig, prs.slides[idx], prs)
        print(f"detected: {detected}")
        if detected != SLIDE_TYPE_QPILL_QTEXT_TABLE_MCQ:
            sys.exit(1)
        shapes = find_shapes(prs.slides[idx])
        q_clone = format_question_text(shapes["question_text_el"])
        ans_clone = format_mcq_answer_text(0, shapes["mcq_text_els"][0])
        print(f"question text normAutofit: {_has_norm_autofit(q_clone)}")
        print(f"answer A normAutofit: {_has_norm_autofit(ans_clone)}")
        out = save_formatted_slide(path, idx)
        print(f"saved: {out}")
    elif len(sys.argv) == 2:
        run_all(sys.argv[1])
    else:
        print("Usage: python -m ...test_formatter <pptx_path> [<slide_index>]")
        sys.exit(1)

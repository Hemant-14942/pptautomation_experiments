"""Save demo outputs for 1, 2, and 3 image layouts from ptest1.

Slide 19 (index 18) has a heading so it is NOT image_only by the detector,
but we can still run the picture formatter on its images to preview the grid.

Run:
    python3 -m app.core.fit_layout.layout_formatter.image_only.tests.demo_outputs
"""

import os
import copy

from pptx import Presentation

from app.core.fit_layout.layout_formatter.slide_type_detector import IMAGE_Y_THRESHOLD_EMU
from app.core.fit_layout.layout_formatter.image_only.formatter import format_image_only_slide

INPUT = "app/data/input/ptest1.pptx"
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "output")


def _content_pictures(slide):
    """Pictures below the header logo area, sorted top then left."""
    items = []
    for sp in slide.shapes:
        stype = sp.shape_type
        if stype is not None and getattr(stype, "name", "") == "PICTURE":
            if sp.top is not None and sp.top > IMAGE_Y_THRESHOLD_EMU:
                items.append((sp.top, sp.left, sp._element, sp.image.size))
    items.sort(key=lambda x: (x[0], x[1]))
    return [x[2] for x in items], [x[3] for x in items]


def _apply_and_save(prs, slide_index, pic_els, pic_sizes, out_name):
    """Replace pictures on one slide and save the deck."""
    slide = prs.slides[slide_index]
    clones = format_image_only_slide(pic_els, pic_sizes)
    spTree = slide.shapes._spTree
    for old in pic_els:
        spTree.remove(old)
    for clone in clones:
        spTree.append(clone)
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    out_path = os.path.join(OUTPUT_DIR, out_name)
    prs.save(out_path)
    return out_path


def main():
    # --- Demo A: real image_only slide (PPT slide 16, index 15) — 1 image ---
    prs1 = Presentation(INPUT)
    pic_els, pic_sizes = _content_pictures(prs1.slides[15])
    path1 = _apply_and_save(
        prs1, 15, pic_els[:1], pic_sizes[:1],
        "ptest1_slide16_1image_formatted.pptx",
    )
    print(f"1-image demo saved: {path1}")

    # --- Demo B: slide 19 pictures, first 2 only — 2-image side by side ---
    prs2 = Presentation(INPUT)
    pic_els, pic_sizes = _content_pictures(prs2.slides[18])
    path2 = _apply_and_save(
        prs2, 18, pic_els[:2], pic_sizes[:2],
        "ptest1_slide19_2image_formatted.pptx",
    )
    print(f"2-image demo saved: {path2}")

    # --- Demo C: slide 19 all 3 pictures — 3-image balanced grid ---
    prs3 = Presentation(INPUT)
    pic_els, pic_sizes = _content_pictures(prs3.slides[18])
    path3 = _apply_and_save(
        prs3, 18, pic_els[:3], pic_sizes[:3],
        "ptest1_slide19_3image_formatted.pptx",
    )
    print(f"3-image demo saved: {path3}")


if __name__ == "__main__":
    main()

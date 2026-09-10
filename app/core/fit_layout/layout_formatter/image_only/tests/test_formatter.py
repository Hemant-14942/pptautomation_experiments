"""Test the image_only formatter.

Run on one slide manually:
    python -m app.core.fit_layout.layout_formatter.image_only.tests.test_formatter <pptx_path> <slide_index>
"""

import os
import sys

from pptx import Presentation
from pptx.util import Emu

from app.core.style_parser import parse_input_slide_signature
from app.core.fit_layout.layout_formatter.slide_type_detector import (
    IMAGE_Y_THRESHOLD_EMU,
    SLIDE_TYPE_IMAGE_ONLY,
    detect_slide_type,
)
from app.core.fit_layout.layout_formatter.image_only.constants import get_image_boxes
from app.core.fit_layout.layout_formatter.image_only.formatter import format_image_only_slide

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "output")

XFRM_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def _get_off_ext(el):
    """Read position and size from a shape's XML."""
    xfrm = el.find(".//p:spPr/a:xfrm", XFRM_NS)
    off = xfrm.find("a:off", XFRM_NS)
    ext = xfrm.find("a:ext", XFRM_NS)
    # Picture XML sometimes stores coords as "6673107.0" — use float() first
    return (
        (int(float(off.get("x"))), int(float(off.get("y")))),
        (int(float(ext.get("cx"))), int(float(ext.get("cy")))),
    )


def _find_content_pictures(slide):
    """Find all content pictures on the slide (below header logo area).

    Returns lists sorted top-to-bottom, then left-to-right so image order
    is stable every time we run the formatter.
    """
    pictures = []
    for sp in slide.shapes:
        stype = sp.shape_type
        if stype is not None and getattr(stype, "name", "") == "PICTURE":
            if sp.top is not None and sp.top > IMAGE_Y_THRESHOLD_EMU:
                pictures.append((sp.top, sp.left, sp._element, sp.image.size))
    # sort by vertical position first, then horizontal
    pictures.sort(key=lambda item: (item[0], item[1]))
    pic_els = [p[2] for p in pictures]
    pic_sizes = [p[3] for p in pictures]
    return pic_els, pic_sizes


def _picture_within_box(pic_off, pic_ext, box):
    """True if the picture fits entirely inside the slot box."""
    return (
        pic_off[0] >= box["x"]
        and pic_off[1] >= box["y"]
        and pic_off[0] + pic_ext[0] <= box["x"] + box["width"]
        and pic_off[1] + pic_ext[1] <= box["y"] + box["height"]
    )


def verify_slide(pptx_path: str, slide_index: int) -> dict:
    """Check detection + formatter for one slide."""
    prs = Presentation(pptx_path)
    slide = prs.slides[slide_index]
    sig = parse_input_slide_signature(pptx_path)[slide_index]

    detected_type = detect_slide_type(sig, slide, prs)
    if detected_type != SLIDE_TYPE_IMAGE_ONLY:
        return {
            "pass": False,
            "reason": f"slide type detected as {detected_type!r}, not image_only",
            "detected_type": detected_type,
        }

    pic_els, pic_sizes = _find_content_pictures(slide)
    if not pic_els:
        return {
            "pass": False,
            "reason": "detector matched but could not locate picture shapes",
            "detected_type": detected_type,
        }

    pic_clones = format_image_only_slide(pic_els, pic_sizes)
    expected_boxes = get_image_boxes(len(pic_els))

    pic_results = []
    all_ok = True
    for pic_clone, pic_size, box in zip(pic_clones, pic_sizes, expected_boxes):
        pic_off, pic_ext = _get_off_ext(pic_clone)
        expected_aspect = pic_size[0] / pic_size[1]
        actual_aspect = pic_ext[0] / pic_ext[1]
        aspect_ok = abs(actual_aspect - expected_aspect) < 0.01
        within_box = _picture_within_box(pic_off, pic_ext, box)
        ok = aspect_ok and within_box
        all_ok = all_ok and ok
        pic_results.append({
            "off": pic_off,
            "ext": pic_ext,
            "aspect_ok": aspect_ok,
            "within_box": within_box,
        })

    return {
        "pass": all_ok,
        "detected_type": detected_type,
        "image_count": len(pic_els),
        "images": pic_results,
    }


def save_formatted_slide(pptx_path: str, slide_index: int, output_dir: str = OUTPUT_DIR) -> str | None:
    """Apply formatter and save a new pptx for visual checking."""
    prs = Presentation(pptx_path)
    slide = prs.slides[slide_index]
    sig = parse_input_slide_signature(pptx_path)[slide_index]

    if detect_slide_type(sig, slide, prs) != SLIDE_TYPE_IMAGE_ONLY:
        return None

    pic_els, pic_sizes = _find_content_pictures(slide)
    if not pic_els:
        return None

    pic_clones = format_image_only_slide(pic_els, pic_sizes)
    spTree = slide.shapes._spTree
    for old_el in pic_els:
        spTree.remove(old_el)
    for clone in pic_clones:
        spTree.append(clone)

    os.makedirs(output_dir, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(pptx_path))[0]
    out_path = os.path.join(output_dir, f"{base_name}_slide{slide_index + 1}_formatted.pptx")
    prs.save(out_path)
    return out_path


def _print_report(pptx_path: str, slide_index: int, report: dict) -> None:
    print(f"\n{pptx_path} | slide index {slide_index} (PPT slide {slide_index + 1})")
    print(f"  detected_type: {report.get('detected_type')}")
    if not report["pass"] and "reason" in report:
        print(f"  RESULT: FAIL — {report['reason']}")
        return

    print(f"  image_count: {report.get('image_count')}")
    for i, img in enumerate(report.get("images", []), start=1):
        off_in = (Emu(img["off"][0]).inches, Emu(img["off"][1]).inches)
        ext_in = (Emu(img["ext"][0]).inches, Emu(img["ext"][1]).inches)
        print(
            f"  image {i}: off={off_in[0]:.2f},{off_in[1]:.2f}in  "
            f"ext={ext_in[0]:.2f}x{ext_in[1]:.2f}in  "
            f"aspect_ok={img['aspect_ok']}  within_box={img['within_box']}"
        )
    print(f"  RESULT: {'PASS' if report['pass'] else 'FAIL'}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python -m ...test_formatter <pptx_path> <slide_index>")
        sys.exit(1)
    path, idx = sys.argv[1], int(sys.argv[2])
    result = verify_slide(path, idx)
    _print_report(path, idx, result)
    if result["pass"]:
        saved_path = save_formatted_slide(path, idx)
        print(f"  saved formatted pptx: {saved_path}")

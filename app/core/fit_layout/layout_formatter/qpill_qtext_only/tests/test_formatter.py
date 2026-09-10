"""Test qpill_qtext_only formatter.

Run one slide:
    python -m app.core.fit_layout.layout_formatter.qpill_qtext_only.tests.test_formatter app/data/input/p.pptx 16

Run all matching slides in a deck:
    python -m app.core.fit_layout.layout_formatter.qpill_qtext_only.tests.test_formatter app/data/input/p.pptx
"""

import os
import sys

from pptx import Presentation

from app.core.style_parser import parse_input_slide_signature
from app.core.fit_layout.layout_formatter.slide_type_detector import (
    SLIDE_TYPE_QPILL_QTEXT_ONLY,
    detect_slide_type,
)
from app.core.fit_layout.layout_formatter.qpill_qtext_only.formatter import (
    format_question_label,
    format_question_pill,
    format_question_text,
)
from app.core.fit_layout.layout_formatter.qpill_shape_utils import find_question_text_element
from app.utils.xml_helpers import prst_geom

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "output")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def _has_norm_autofit(el) -> bool:
    txBody = el.find(".//p:txBody", NS)
    if txBody is None:
        return False
    bodyPr = txBody.find("a:bodyPr", NS)
    return bodyPr is not None and bodyPr.find("a:normAutofit", NS) is not None


def find_shapes(slide) -> dict | None:
    question_pill_el = None
    question_label_el = None
    question_text_el = None

    for sp in slide.shapes:
        elem = sp._element
        geom = prst_geom(elem)
        if geom == "roundRect" and question_pill_el is None:
            question_pill_el = elem
            continue

        if not sp.has_text_frame:
            continue
        text = sp.text_frame.text.strip()
        if not text:
            continue

        if text == "Question" and question_label_el is None:
            question_label_el = elem

    question_text_el = find_question_text_element(slide, use_mcq_ceiling=False)

    if (
        question_pill_el is None
        or question_label_el is None
        or question_text_el is None
    ):
        return None

    return {
        "question_pill_el": question_pill_el,
        "question_label_el": question_label_el,
        "question_text_el": question_text_el,
    }


def save_formatted_slide(pptx_path: str, slide_index: int, output_dir: str = OUTPUT_DIR) -> str | None:
    prs = Presentation(pptx_path)
    slide = prs.slides[slide_index]
    sig = parse_input_slide_signature(pptx_path)[slide_index]

    if detect_slide_type(sig, slide, prs) != SLIDE_TYPE_QPILL_QTEXT_ONLY:
        return None

    shapes = find_shapes(slide)
    if shapes is None:
        return None

    spTree = slide.shapes._spTree

    def reposition(orig_el, clone_el):
        if orig_el.getparent() is not None:
            spTree.remove(orig_el)
        spTree.append(clone_el)

    reposition(shapes["question_pill_el"], format_question_pill(shapes["question_pill_el"]))
    reposition(shapes["question_label_el"], format_question_label(shapes["question_label_el"]))
    reposition(shapes["question_text_el"], format_question_text(shapes["question_text_el"]))

    os.makedirs(output_dir, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(pptx_path))[0]
    out_path = os.path.join(output_dir, f"{base_name}_slide{slide_index + 1}_formatted.pptx")
    prs.save(out_path)
    return out_path


def run_all(pptx_path: str) -> None:
    prs = Presentation(pptx_path)
    signatures = parse_input_slide_signature(pptx_path)

    matched = 0
    saved = 0
    for idx, slide in enumerate(prs.slides):
        if detect_slide_type(signatures[idx], slide, prs) != SLIDE_TYPE_QPILL_QTEXT_ONLY:
            continue
        matched += 1
        out_path = save_formatted_slide(pptx_path, idx)
        status = out_path if out_path else "SKIPPED (shapes not found)"
        print(f"  slide index {idx:2d} (PPT slide {idx + 1:2d}): {status}")
        if out_path:
            saved += 1

    print(f"\n{matched} slides detected as qpill_qtext_only, {saved} formatted and saved to {OUTPUT_DIR}")


if __name__ == "__main__":
    if len(sys.argv) == 3:
        path, idx = sys.argv[1], int(sys.argv[2])
        prs = Presentation(path)
        sig = parse_input_slide_signature(path)[idx]
        detected = detect_slide_type(sig, prs.slides[idx], prs)
        print(f"detected: {detected}")
        if detected != SLIDE_TYPE_QPILL_QTEXT_ONLY:
            sys.exit(1)
        shapes = find_shapes(prs.slides[idx])
        q_clone = format_question_text(shapes["question_text_el"])
        print(f"question text normAutofit: {_has_norm_autofit(q_clone)}")
        out = save_formatted_slide(path, idx)
        print(f"saved: {out}")
    elif len(sys.argv) == 2:
        run_all(sys.argv[1])
    else:
        print("Usage: python -m ...test_formatter <pptx_path> [<slide_index>]")
        sys.exit(1)
